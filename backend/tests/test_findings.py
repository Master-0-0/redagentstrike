import copy
import hashlib
import json
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from workbench.api import create_app
from workbench.findings import FindingsStore, ArchiveError
from workbench.netstore import HttpEvent


def manifest(key='test', body=b'print("fixture")'):
    return dict(submission_key=key, op='create', finding=dict(title='测试成果', target='local fixture', summary='已保存材料', severity='unknown', preconditions='fixture only', steps=[dict(id='step-1', description='查看测试材料', artifact_client_ids=['poc'])], expected_result='fixture', actual_result='fixture', agent_assessment='reproduced'), source=dict(source_type='manual'), artifacts=[dict(client_id='poc', kind='poc', filename='fixture.py', size=len(body), sha256=hashlib.sha256(body).hexdigest(), media_type='text/plain', step_ids=['step-1'])], evidence_gaps=[dict(category=c, status='missing', reason='未采集') for c in ('http', 'screenshot', 'tool_output')])


def prepare(client, m):
    response = client.post('/finding-submissions', json=m)
    assert response.status_code == 200, response.text
    return response.json()


def upload(client, s, body=b'print("fixture")'):
    response = client.put(f"/finding-submissions/{s['id']}/artifacts/poc", content=body, headers={'content-type': 'text/plain'})
    assert response.status_code == 200, response.text


def commit(client, s):
    response = client.post(f"/finding-submissions/{s['id']}/commit")
    assert response.status_code == 200, response.text
    return response.json()


def test_immutable_archive_retry_restart_and_missing(client, settings):
    m = manifest()
    s = prepare(client, m)
    assert client.get('/findings').json()['total'] == 0
    assert client.post(f"/finding-submissions/{s['id']}/commit").status_code == 409
    upload(client, s)
    receipt = commit(client, s)
    assert receipt['review_status'] == 'pending'
    assert commit(client, s) == receipt
    assert prepare(client, m)['receipt'] == receipt
    upload(client, s)
    fid = receipt['finding_id']
    detail = client.get('/findings/' + fid).json()
    a = detail['artifacts'][0]
    assert detail['steps'][0]['artifact_ids'] == [a['id']]
    url = f"/findings/{fid}/artifacts/{a['id']}"
    assert client.get(url).content == b'print("fixture")'
    restarted = FindingsStore(settings, client.app.state.wb.net)
    assert restarted.detail(fid)['artifacts'][0]['integrity'] == 'ok'
    restarted.close()
    (settings.data_dir / 'findings' / fid / s['id'] / a['id']).unlink()
    assert client.get('/findings/' + fid).json()['artifacts'][0]['integrity'] == 'error'
    assert client.get(url).status_code == 409
    assert client.put(f"/finding-submissions/{s['id']}/artifacts/poc", content=b'print("fixture")', headers={'content-type':'text/plain'}).status_code == 409


def test_conflicts_append_and_human_review(client):
    m = manifest()
    s = prepare(client, m)
    upload(client, s)
    fid = commit(client, s)['finding_id']
    edited = client.patch('/findings/' + fid, json=dict(version=1, changes=dict(review_status='confirmed'), note='人工核实'))
    assert edited.status_code == 200
    assert client.patch('/findings/' + fid, json=dict(version=1, changes=dict(title='stale'))).status_code == 409
    m['finding']['summary'] = 'different'
    assert client.post('/finding-submissions', json=m).json()['error_code'] == 'idempotency_conflict'
    gid = edited.json()['evidence_gaps'][0]['id']
    append = dict(submission_key='append', op='append', finding_id=fid, source=dict(source_type='cairn', workspace_path='/fixture', session_id='child', owner_resolution='unresolved'), note='补充人工说明', resolved_gap_ids=[gid])
    receipt = commit(client, prepare(client, append))
    assert receipt['review_status'] == 'confirmed'
    detail = client.get('/findings/' + fid).json()
    assert detail['version'] == 3 and len(detail['artifacts']) == 1
    assert detail['evidence_gaps'][0]['resolved_at']
    assert len(detail['history']) == 3
    assert client.get('/findings', params={'session_id': 'child'}).json()['total'] == 1
    assert client.get('/findings', params={'review_status': 'pending'}).json()['total'] == 0


def test_concurrent_idempotence_and_appends(client):
    m = manifest()
    with ThreadPoolExecutor(max_workers=6) as pool:
        submissions = list(pool.map(lambda _: prepare(client, m), range(6)))
    assert len({s['id'] for s in submissions}) == 1
    s = submissions[0]
    upload(client, s)
    with ThreadPoolExecutor(max_workers=6) as pool:
        receipts = list(pool.map(lambda _: commit(client, s), range(6)))
    assert len({r['finding_id'] for r in receipts}) == 1
    fid = receipts[0]['finding_id']
    def append(i):
        return commit(client, prepare(client, dict(op='append', finding_id=fid, submission_key=f'append-{i}', source=dict(source_type='manual'), note=f'note {i}')))
    with ThreadPoolExecutor(max_workers=6) as pool:
        list(pool.map(append, range(6)))
    detail = client.get('/findings/' + fid).json()
    assert detail['version'] == 7 and len(detail['history']) == 7


def test_crash_after_move_recovery(client, settings):
    s = prepare(client, manifest())
    upload(client, s)
    store = client.app.state.wb.findings
    with patch.object(store, '_commit_database', side_effect=OSError('simulated crash')):
        assert client.post(f"/finding-submissions/{s['id']}/commit").status_code == 503
    assert client.get('/findings').json()['total'] == 0
    assert (settings.data_dir / 'findings' / s['finding_id'] / s['id']).is_dir()
    restarted = FindingsStore(settings, client.app.state.wb.net)
    receipt = restarted.commit(s['id'])
    assert restarted.commit(s['id']) == receipt
    assert restarted.listing({}, 1, 20)['total'] == 1
    restarted.close()


def test_network_snapshot_survives_source_deletion(client):
    event = HttpEvent(id='event-1', run_id='run-1', ts='2026-09-24T00:00:00Z', host='fixture', method='GET', url='http://fixture/', status=200, summary='fixture', req_headers={'Authorization': 'Bearer sensitive'}, req_body='', resp_headers={}, resp_body='response bytes', evidence_id=None, failed=False)
    client.app.state.wb.net.insert(event)
    m = manifest()
    m['network_refs'] = [dict(event_id='event-1', run_id='wrong')]
    assert client.post('/finding-submissions', json=m).json()['error_code'] == 'network_run_mismatch'
    m['submission_key'] = 'valid-network'
    m['network_refs'][0]['run_id'] = 'run-1'
    s = prepare(client, m)
    client.app.state.wb.net._conn.execute('DELETE FROM net_http')
    upload(client, s)
    fid = commit(client, s)['finding_id']
    a = client.get('/findings/' + fid).json()['artifacts'][1]
    assert a['fidelity'] == 'transformed' and a['truncation'] == 'unknown'
    url = f"/findings/{fid}/artifacts/{a['id']}"
    assert client.get(url).json()['resp_body'] == 'response bytes'
    assert 'sensitive' in client.get(url + '?preview=true').json()['text']
    assert client.get('/findings?run_id=run-1').json()['total'] == 1


def test_invalid_refs_types_limits_and_inert_preview(client):
    m = manifest(body=b'<script>alert(1)</script>')
    m['artifacts'][0]['media_type'] = 'text/html'
    s = prepare(client, m)
    response = client.put(f"/finding-submissions/{s['id']}/artifacts/poc", content=b'<script>alert(1)</script>', headers={'content-type': 'text/html'})
    assert response.status_code == 200
    fid = commit(client, s)['finding_id']
    aid = client.get('/findings/' + fid).json()['artifacts'][0]['id']
    url = f'/findings/{fid}/artifacts/{aid}'
    assert client.get(url).headers['content-type'] == 'application/octet-stream'
    assert client.get(url + '?preview=true').headers['content-type'] == 'application/json'
    for mutate in [lambda x: x['artifacts'][0].update(client_id='../escape'), lambda x: x['artifacts'][0].update(fidelity='raw_capture'), lambda x: x['artifacts'][0].update(step_ids=['absent']), lambda x: x.update(source={'source_type':'cairn'}), lambda x: x['finding'].update(review_status='confirmed')]:
        bad = manifest('bad')
        mutate(bad)
        assert client.post('/finding-submissions', json=bad).status_code in (400, 422)
    m = manifest('large'); m['artifacts'][0]['size'] = 21 * 1024 * 1024
    assert client.post('/finding-submissions', json=m).status_code == 413
    assert client.post('/finding-submissions', content=b'x' * (1024 * 1024 + 1)).status_code == 413
    s = prepare(client, manifest('hash'))
    assert client.put(f"/finding-submissions/{s['id']}/artifacts/poc", content=b'x', headers={'content-type':'text/plain'}).json()['error_code'] == 'hash_mismatch'
    assert client.get('/findings?page_size=101').status_code == 422


def test_corrupt_database_isolated(settings):
    settings.data_dir.mkdir(exist_ok=True)
    path = settings.data_dir / 'findings.sqlite'
    path.write_bytes(b'corrupt database, must never be deleted')
    with TestClient(create_app(settings)) as client:
        assert client.get('/findings').status_code == 503
        assert client.get('/health').status_code == 200
        assert path.read_bytes() == b'corrupt database, must never be deleted'


def test_preview_limit_and_images(client):
    body = b'Authorization: Bearer secret\n' + b'A' * 70000
    s = prepare(client, manifest('preview', body))
    upload(client, s, body)
    fid = commit(client, s)['finding_id']
    aid = client.get('/findings/' + fid).json()['artifacts'][0]['id']
    url = f'/findings/{fid}/artifacts/{aid}'
    assert client.get(url).content == body
    preview = client.get(url + '?preview=true').json()
    assert preview['truncated'] and 'Bearer secret' in preview['text']
    from PIL import Image
    out = BytesIO(); Image.new('RGB', (2, 2)).save(out, format='PNG')
    data = out.getvalue()
    m = manifest('image', data); m['evidence_gaps'].append(dict(category='poc', status='not_applicable', reason='image fixture')); m['artifacts'][0].update(kind='response_screenshot', media_type='image/png')
    s = prepare(client, m)
    assert client.put(f"/finding-submissions/{s['id']}/artifacts/poc", content=data, headers={'content-type':'image/png'}).status_code == 200
    fid = commit(client, s)['finding_id']
    aid = client.get('/findings/' + fid).json()['artifacts'][0]['id']
    assert client.get(f'/findings/{fid}/artifacts/{aid}?preview=true').headers['content-type'] == 'image/png'


def test_explicit_gaps_and_upload_recovery(client):
    m = manifest('gaps'); m['evidence_gaps'] = []
    assert client.post('/finding-submissions', json=m).json()['error_code'] == 'invalid_manifest'
    s = prepare(client, manifest('recovery'))
    response = client.put(f"/finding-submissions/{s['id']}/artifacts/poc", content=b'bad', headers={'content-type':'text/plain'})
    assert response.status_code == 409
    progress = client.get('/finding-submissions/' + s['id']).json()
    assert progress['status'] == 'failed' and progress['error']['retryable']
    upload(client, s)
    assert commit(client, s)['ok']


def test_database_transaction_rollback_after_move(client):
    s = prepare(client, manifest())
    upload(client, s)
    store = client.app.state.wb.findings
    original = store.insert
    def failing_insert(table, value, fid):
        original(table, value, fid)
        if table == 'history':
            raise OSError('disk full at transaction boundary')
    with patch.object(store, 'insert', side_effect=failing_insert):
        assert client.post(f"/finding-submissions/{s['id']}/commit").status_code == 503
    assert client.get('/findings').json()['total'] == 0
    progress = client.get('/finding-submissions/' + s['id']).json()
    assert progress['status'] == 'failed'
    assert progress['last_completed_stage'] == 'files_moved'
    fid = commit(client, s)['finding_id']
    detail = client.get('/findings/' + fid).json()
    assert len(detail['sources']) == 1 and len(detail['history']) == 1


def test_missing_network_and_invalid_image(client):
    m = manifest('missing-net'); m['network_refs'] = [dict(event_id='absent',run_id='not-current')]
    assert client.post('/finding-submissions', json=m).json()['error_code'] == 'network_event_missing'
    m = manifest('fake-image'); m['artifacts'][0]['media_type'] = 'image/png'
    s = prepare(client, m)
    response = client.put(f"/finding-submissions/{s['id']}/artifacts/poc",content=b'print("fixture")',headers={'content-type':'image/png'})
    assert response.status_code == 400
    assert client.get('/findings').json()['total'] == 0


def test_pagination_and_field_edit_history(client):
    for index in range(23):
        m = manifest(f'page-{index}'); m['artifacts']=[]; m['finding']['steps'][0]['artifact_client_ids']=[]
        m['evidence_gaps'].append(dict(category='poc',status='missing',reason='fixture'))
        commit(client, prepare(client,m))
    first = client.get('/findings').json()
    second = client.get('/findings?page=2').json()
    assert first['total'] == 23 and len(first['items']) == 20 and len(second['items']) == 3
    assert not set(f['id'] for f in first['items']) & set(f['id'] for f in second['items'])
    fid=first['items'][0]['id']
    edit=client.patch('/findings/'+fid,json=dict(version=1,changes=dict(steps=[dict(id='updated-step',description='New version')],summary='Edited summary'),note='manual edit'))
    assert edit.status_code == 200
    assert edit.json()['history'][-1]['before']['steps'][0]['id'] == 'step-1'
    for invalid in ('not-an-array', [42], [dict(id='fake', description='fake', artifact_ids=['fake-id'])]):
        assert client.patch('/findings/'+fid,json=dict(version=2,changes=dict(steps=invalid))).status_code == 400


def test_network_move_before_stage_update_recovers_without_event(client, settings):
    event = HttpEvent(id='crash-network', run_id='fixture-run', ts='2026-09-24T00:00:00Z', host='fixture', method='GET', url='http://fixture', status=200, summary='fixture', req_headers={}, req_body='', resp_headers={}, resp_body='durable snapshot', evidence_id=None, failed=False)
    client.app.state.wb.net.insert(event)
    m=manifest('network-crash');m['network_refs']=[dict(event_id=event.id,run_id=event.run_id)]
    # The filesystem rename succeeds, then the directory flush fails. Planned
    # bytes/hash must already be in SQLite; the next attempt needs no NetStore.
    with patch('workbench.findings.sync_dir',side_effect=OSError('crash after network rename')):
        assert client.post('/finding-submissions',json=m).status_code==503
    client.app.state.wb.net._conn.execute('DELETE FROM net_http')
    s=prepare(client,m);upload(client,s)
    fid=commit(client,s)['finding_id']
    detail=client.get('/findings/'+fid).json()
    assert detail['artifacts'][1]['integrity']=='ok'


def test_preview_preserves_all_saved_credential_fields(tmp_path):
    from workbench.findings_api import read_preview

    text = ('Authorization: Bearer fixture-token\n'
            'Proxy-Authorization: Basic fixture-proxy\n'
            'Cookie: sid=fixture-cookie\nSet-Cookie: sid=fixture-cookie; HttpOnly\n'
            '{"password":"fixture-password","passwd":"fixture-passwd",'
            '"token":"fixture-token","secret":"fixture-secret","api_key":"fixture-key"}\n'
            'https://fixture.invalid/?token=fixture-token&password=fixture-password&secret=fixture-secret&api_key=fixture-key\n'
            '中文证据保持原样\n')
    path = tmp_path / 'evidence.txt'
    path.write_text(text, encoding='utf-8')
    assert read_preview(path) == {'text': text, 'truncated': False}


def test_export_selected_reports_and_validation(client):
    import zipfile
    from docx import Document

    first = prepare(client, manifest('export-a'))
    upload(client, first)
    fid = commit(client, first)['finding_id']
    other = manifest('export-b')
    other['finding']['title'] = '不应包含的未选成果'
    second = prepare(client, other)
    upload(client, second)
    commit(client, second)
    before = client.get('/findings/' + fid).json()
    response = client.post('/findings/export', json={'finding_ids': [fid, fid], 'format': 'docx'})
    assert response.status_code == 200
    assert 'attachment;' in response.headers['content-disposition']
    doc = Document(BytesIO(response.content))
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert '成果数量：1' in text and '测试成果' in text and '待复核' in text
    assert '不应包含的未选成果' not in text
    assert before['artifacts'][0]['sha256'] in text
    pdf = client.post('/findings/export', json={'finding_ids': [fid], 'format': 'pdf'})
    assert pdf.status_code == 200 and pdf.content.startswith(b'%PDF-')
    assert pdf.headers['content-type'] == 'application/pdf'
    assert client.get('/findings/' + fid).json() == before
    for data in ({'finding_ids': [], 'format': 'pdf'}, {'finding_ids': [fid], 'format': 'html'}, {'finding_ids': [fid] * 101, 'format': 'docx'}, {'finding_ids': [None], 'format': 'pdf'}):
        assert client.post('/findings/export', json=data).status_code == 422
    assert client.post('/findings/export', json={'finding_ids': ['missing'], 'format': 'docx'}).status_code == 404


def test_export_includes_saved_text_and_images(client):
    from PIL import Image
    from docx import Document
    import zipfile

    body = b'Authorization: Bearer report-secret\n  saved evidence <script>inert</script>\n'
    m = manifest('report-evidence', body)
    image = BytesIO()
    Image.new('RGB', (600, 160), 'lightblue').save(image, format='PNG')
    png = image.getvalue()
    m['artifacts'].append(dict(client_id='image', kind='browser_screenshot', filename='evidence.png', size=len(png), sha256=hashlib.sha256(png).hexdigest(), media_type='image/png', step_ids=['step-1']))
    s = prepare(client, m)
    upload(client, s, body)
    assert client.put(f"/finding-submissions/{s['id']}/artifacts/image", content=png, headers={'content-type': 'image/png'}).status_code == 200
    fid = commit(client, s)['finding_id']
    response = client.post('/findings/export', json={'finding_ids': [fid], 'format': 'docx'})
    assert response.status_code == 200
    doc = Document(BytesIO(response.content))
    assert body.decode().strip() in '\n'.join(p.text for p in doc.paragraphs)
    assert len(doc.inline_shapes) == 1
    with zipfile.ZipFile(BytesIO(response.content)) as z:
        assert png in [z.read(n) for n in z.namelist() if n.startswith('word/media/')]
    pdf = client.post('/findings/export', json={'finding_ids': [fid], 'format': 'pdf'})
    assert pdf.status_code == 200 and b'/Subtype /Image' in pdf.content
